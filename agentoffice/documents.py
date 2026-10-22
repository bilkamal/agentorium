"""Text extraction for uploaded documents (txt/md, docx, xlsx, pdf)."""
from __future__ import annotations

from pathlib import Path

ALLOWED_EXTENSIONS = {".txt", ".md", ".docx", ".xlsx", ".pdf"}
MAX_UPLOAD_BYTES = 25 * 1024 * 1024


class DocumentError(Exception):
    pass


def extract_text(path: Path, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentError(f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}")
    try:
        if ext in (".txt", ".md"):
            return path.read_text(encoding="utf-8", errors="replace")
        if ext == ".docx":
            return _read_docx(path)
        if ext == ".xlsx":
            return _read_xlsx(path)
        if ext == ".pdf":
            return _read_pdf(path)
    except DocumentError:
        raise
    except Exception as exc:  # corrupted / password-protected files etc.
        raise DocumentError(f"Could not read {filename}: {exc}") from exc
    raise DocumentError("Unreachable")


def _read_docx(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    # Tables are flattened as pipe-separated rows so the model keeps the structure.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            if row is None or all(v is None for v in row):
                continue
            parts.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentError("PDF is password protected") from exc
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"--- page {i} ---\n{text}")
    return "\n".join(pages)
